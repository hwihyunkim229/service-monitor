from datetime import datetime
from pathlib import Path
import threading
import time

from openpyxl import load_workbook

from app.database import SessionLocal
from app.models.lot_process_rev import LotProcessRev

_CACHE_SECONDS = 30
_cache = {}
_cache_loaded_at = 0.0
_cache_lock = threading.Lock()


def _clean(value):
    return str(value).strip() if value is not None else ""


def _format_rev(value):
    value = _clean(value)
    if not value:
        return ""
    compact = value.lower().replace(" ", "").replace(".", "")
    if compact.startswith("rev"):
        value = compact[3:]
    return f"Rev{value}"


def read_process_rev_workbook(source):
    """Read either the simple upload sheet or the original matrix workbook."""
    workbook = load_workbook(source, data_only=True, read_only=True)
    mapping = {}
    conflicts = []
    try:
        simple_sheet = next(
            (
                sheet
                for sheet in workbook.worksheets
                if _clean(sheet.cell(1, 1).value).upper() == "LOT NO."
                and _clean(sheet.cell(1, 2).value) in {"출고 Rev", "공정 Rev"}
            ),
            None,
        )
        if simple_sheet is not None:
            for row, values in enumerate(
                simple_sheet.iter_rows(min_row=2, max_col=2, values_only=True),
                start=2,
            ):
                lot = _clean(values[0]).upper()
                process_rev = _format_rev(values[1])
                if not lot and not process_rev:
                    continue
                if not lot or not process_rev:
                    raise ValueError(
                        f"{row}행의 LOT No. 또는 출고 Rev가 비어 있습니다."
                    )
                previous = mapping.get(lot)
                current = {"process_rev": process_rev}
                if previous and previous != current:
                    conflicts.append(lot)
                mapping[lot] = current
            if conflicts:
                examples = ", ".join(sorted(set(conflicts))[:5])
                raise ValueError(f"서로 다른 Rev에 중복된 LOT가 있습니다: {examples}")
            if not mapping:
                raise ValueError("'업로드 양식' 시트에 LOT를 한 건 이상 입력하세요.")
            return mapping

        worksheet = workbook.active
        rows = list(worksheet.iter_rows(values_only=True))
        if len(rows) < 3 or _clean(rows[1][0]) != "출고 Rev":
            raise ValueError(
                "첫 행에 'LOT No.'와 '출고 Rev' 열이 있는 양식을 사용하세요."
            )
        if _clean(rows[2][0]).upper() != "LOT NO.":
            raise ValueError("3행 첫 셀에서 'LOT No.'를 찾을 수 없습니다.")

        for column in range(1, len(rows[1])):
            process_rev = _format_rev(rows[1][column])
            ring_rev = _clean(rows[0][column])
            if not process_rev:
                continue
            for values in rows[2:]:
                lot = _clean(values[column] if column < len(values) else None).upper()
                if not lot:
                    continue
                previous = mapping.get(lot)
                current = {"process_rev": process_rev, "ring_rev": ring_rev or None}
                if previous and previous != current:
                    conflicts.append(lot)
                mapping[lot] = current
    finally:
        workbook.close()

    if conflicts:
        examples = ", ".join(sorted(set(conflicts))[:5])
        raise ValueError(f"서로 다른 Rev에 중복된 LOT가 있습니다: {examples}")
    if not mapping:
        raise ValueError("업로드 파일에서 LOT 정보를 찾을 수 없습니다.")
    return mapping


def invalidate_process_rev_cache():
    global _cache_loaded_at
    with _cache_lock:
        _cache.clear()
        _cache_loaded_at = 0.0


def import_process_rev_workbook(source, db=None):
    mapping = read_process_rev_workbook(source)
    owns_session = db is None
    db = db or SessionLocal()
    inserted = updated = unchanged = 0
    try:
        existing = {
            item.lot: item
            for item in db.query(LotProcessRev)
            .filter(LotProcessRev.lot.in_(list(mapping)))
            .all()
        }
        now = datetime.now()
        for lot, values in mapping.items():
            item = existing.get(lot)
            if item is None:
                db.add(LotProcessRev(lot=lot, updated_at=now, **values))
                inserted += 1
            elif item.process_rev != values["process_rev"] or (
                "ring_rev" in values and item.ring_rev != values["ring_rev"]
            ):
                item.process_rev = values["process_rev"]
                if "ring_rev" in values:
                    item.ring_rev = values["ring_rev"]
                item.updated_at = now
                updated += 1
            else:
                unchanged += 1
        db.commit()
        invalidate_process_rev_cache()
        return {
            "total": len(mapping),
            "inserted": inserted,
            "updated": updated,
            "unchanged": unchanged,
        }
    except Exception:
        db.rollback()
        raise
    finally:
        if owns_session:
            db.close()


def seed_process_rev_master(path=None):
    source = Path(path or "app/data/process_rev_master.xlsx")
    if not source.exists():
        return None
    db = SessionLocal()
    try:
        if db.query(LotProcessRev.lot).first():
            return None
        return import_process_rev_workbook(source, db)
    finally:
        db.close()


def get_process_rev(lot):
    global _cache_loaded_at
    lot = _clean(lot).upper()
    if not lot:
        return ""
    now = time.monotonic()
    if now - _cache_loaded_at >= _CACHE_SECONDS:
        with _cache_lock:
            if now - _cache_loaded_at >= _CACHE_SECONDS:
                db = SessionLocal()
                try:
                    _cache.clear()
                    _cache.update(
                        db.query(LotProcessRev.lot, LotProcessRev.process_rev).all()
                    )
                    _cache_loaded_at = now
                finally:
                    db.close()
    return _cache.get(lot, "")