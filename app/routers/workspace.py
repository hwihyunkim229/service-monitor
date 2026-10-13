from datetime import date, datetime
from collections import Counter
from io import BytesIO
import json
import re
from fastapi import APIRouter, File, HTTPException, Request, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import String, asc, cast, desc, func, or_
from app.database import SessionLocal
from app.models.cs import CSCase
from app.models.cx import CXAnalysis
from app.models.lab import LabAnalysis
from app.models.quality import QualityInspection
from app.models.quality_memo import QualityCellMemo
from app.models.return_transfer import ReturnTransferItem, ReturnTransferRequest
from app.services.cx_categories import rev_from_serial
from app.services.quality_schema import QUALITY_FIELDS, QUALITY_HEADER_FIELDS
from app.utils.downloads import dated_excel_disposition
from app.utils.query_filters import BLANK_FILTER_VALUE, apply_column_filters
from app.services.products import is_bp_product
from app.utils.dates import calendar_week

router = APIRouter()
SECTIONS = {"cx", "exchange", "quality", "lab", "result"}
MAX_UPLOAD_BYTES = 10 * 1024 * 1024

SECTION_HEADERS = {
    "cx": [
        "월",
        "주",
        "일자",
        "병원명",
        "REV 구분",
        "식별코드",
        "CX",
        "RnD",
        "RnD Team",
        "접수 유형(대)",
        "접수 유형(소)",
        "문의사항",
        "교품 여부",
        "S/N",
        "LOT",
        "Size",
        "접수",
        "처리",
        "처리시간",
        "처리 결과",
        "데이터 손실",
        "손실 건수",
        "불량원인 1",
        "Depth 1-1",
        "Depth 1-2",
        "불량원인 2",
        "Depth 2-1",
        "Depth 2-2",
    ],
    "exchange": [
        "월",
        "주",
        "식별코드",
        "접수일",
        "회수일",
        "병원명",
        "유형(대)",
        "유형(소)",
        "교품",
        "고장 S/N",
        "LOT",
        "Size",
        "리드타임",
        "교체 S/N",
    ],
    "quality": [
        "접수일",
        "회수일",
        "회수품 이관일",
        "병원명",
        "식별코드",
        "S/N",
        "공정 Rev",
        "HW Rev",
        "유형(대)",
        "유형(소)",
        "문의사항",
        "CX 원인1",
        "CX Depth1-1",
        "CX Depth1-2",
        "CX 원인2",
        "CX Depth2-1",
        "CX Depth2-2",
        *QUALITY_HEADER_FIELDS.keys(),
        "품질 분석 완료일",
    ],
    "lab": [
        "식별코드",
        "S/N",
        "품질 분석 완료일",
        "불량 원인",
        "분석 결과",
        "불량 유형",
        "비고",
        "분석 완료일",
        "첨부 자료",
    ],
    "result": [
        "식별코드",
        "접수일",
        "병원명",
        "유형(대)",
        "유형(소)",
        "문의사항",
        "S/N",
        "LOT",
        "Size",
        "교품",
        "CX 처리 결과",
        "품질 분석",
        "연구소 분석",
        "접수 → 회수",
        "회수 → 품질",
        "품질 → 연구소",
        "품질 결과",
        "연구소 결과",
        "완료일",
    ],
}


def clean_text(value):
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    value = str(value).strip()
    return (
        None
        if not value or value == "-" or value.lower() in {"none", "null", "nan"}
        else value
    )


def display_transfer_no(value):
    value = clean_text(value)
    return "N/A" if value and value.startswith("__NA__") else (value or "")


def parse_datetime(value, *, compact=False, default_year=None):
    if value is None or value == "" or value == "-":
        return None
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime(value.year, value.month, value.day)
    text = clean_text(value)
    if not text:
        return None
    try:
        return datetime.fromisoformat(text.replace("Z", ""))
    except ValueError:
        pass
    for fmt in (
        "%Y/%m/%d %H:%M",
        "%Y-%m-%d %H:%M",
        "%Y/%m/%d",
        "%Y-%m-%d",
        "%m/%d/%y %H:%M",
        "%m/%d/%Y %H:%M",
        "%m/%d/%y",
        "%m/%d/%Y",
    ):
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            continue
    korean_date = re.fullmatch(r"(\d{1,2})\s*월\s*(\d{1,2})\s*일", text)
    if korean_date:
        month, day = map(int, korean_date.groups())
        return datetime(int(default_year or datetime.now().year), month, day)
    if compact:
        match = re.fullmatch(r"(\d{1,2})[./-](\d{1,2})\s+(\d{1,2}):(\d{2})", text)
        if match:
            month, day, hour, minute = map(int, match.groups())
            return datetime(int(default_year or datetime.now().year), month, day, hour, minute)
    raise ValueError(f"날짜 형식을 확인해 주세요: {text}")


def fmt_date(value, separator="-"):
    return value.strftime(f"%Y{separator}%m{separator}%d") if value else ""


def fmt_compact(value):
    return f"{value.month}/{value.day} {value.strftime('%H:%M')}" if value else ""


def handling_minutes(received_at, processed_at):
    if not received_at or not processed_at:
        return "-"
    minutes = round((processed_at - received_at).total_seconds() / 60)
    return f"{minutes}분" if minutes >= 0 else "-"


def lead_days(start, end):
    if not start or not end:
        return "-"
    start = start.date() if isinstance(start, datetime) else start
    end = end.date() if isinstance(end, datetime) else end
    return f"{(end - start).days}일"


def quality_status(cs, quality):
    if cs.return_at and not quality:
        return "대기 중"
    if quality and not (
        quality.quality_result in {"OK", "NG"} and quality.lab_transfer_date
    ):
        return "분석 중"
    if quality and quality.quality_result in {"OK", "NG"} and quality.lab_transfer_date:
        return "분석 완료"
    return "-"


def lab_status(quality, lab):
    if quality and quality.lab_transfer_date and not lab:
        return "대기 중"
    if lab and not lab.analysis_completed_date:
        return "분석 중"
    if lab and lab.analysis_completed_date:
        return "분석 완료"
    return "-"


def resolved_quality_rev(cs, quality):
    if quality and quality.rev:
        return quality.rev
    lot = (cs.serial_no or "").split("-")[0]
    try:
        from app.routers.quality import get_rev

        return get_rev(lot)
    except Exception:
        return cs.rev or ""


def page_rows(db, request, product, section):
    params = request.query_params
    if section == "cx":
        query = db.query(CSCase).filter(CSCase.product == product)
        try:
            selected_year = int(params.get("year") or datetime.now().year)
        except (TypeError, ValueError):
            selected_year = datetime.now().year
        query = query.filter(
            func.strftime("%Y", CSCase.date_only) == str(selected_year)
        )
        keyword = (params.get("keyword") or "").strip()
        if keyword:
            token = f"%{keyword}%"
            query = query.filter(
                or_(
                    CSCase.transfer_no.ilike(token),
                    CSCase.hospital_name.ilike(token),
                    CSCase.serial_no.ilike(token),
                    CSCase.description.ilike(token),
                )
            )
        if params.get("review") == "1":
            query = query.filter(
                CSCase.category_main == "기타", CSCase.category_sub == "기타"
            )
        week_filter = (params.get("week") or "").strip()
        if week_filter:
            try:
                week_number = int(week_filter)
            except ValueError:
                week_number = 0
            query = query.filter(
                CSCase.week == week_number
                if 1 <= week_number <= 53
                else CSCase.id == -1
            )
        cases = query.order_by(desc(CSCase.date_only), desc(CSCase.id)).all()
        codes = [row.transfer_no for row in cases if row.transfer_no]
        analyses = (
            {
                row.transfer_no: row
                for row in db.query(CXAnalysis)
                .filter(CXAnalysis.transfer_no.in_(codes))
                .all()
            }
            if codes
            else {}
        )
        return [(row, analyses.get(row.transfer_no)) for row in cases]

    if section == "exchange":
        effective_received_at = func.coalesce(CSCase.date_only, CSCase.received_at)
        query = (
            db.query(CSCase, CXAnalysis)
            .outerjoin(CXAnalysis, CSCase.transfer_no == CXAnalysis.transfer_no)
            .filter(
                CSCase.product == product,
                func.lower(func.trim(CSCase.exchange_type)).in_(["o", "cradle"]),
            )
        )
        keyword = (params.get("keyword") or "").strip()
        if keyword:
            token = f"%{keyword}%"
            search_columns = [
                CSCase.transfer_no,
                effective_received_at,
                CSCase.return_at,
                CSCase.week,
                CSCase.hospital_name,
                CSCase.category_main,
                CSCase.category_sub,
                CSCase.exchange_type,
                CSCase.serial_no,
                CSCase.replacement_serial_no,
                CSCase.lot,
                CSCase.size,
                CXAnalysis.issue1,
                CXAnalysis.depth1_1,
                CXAnalysis.depth1_2,
                CXAnalysis.issue2,
                CXAnalysis.depth2_1,
                CXAnalysis.depth2_2,
                CXAnalysis.result_summary,
            ]
            query = query.filter(
                or_(*[cast(column, String).ilike(token) for column in search_columns])
            )
        query = apply_column_filters(
            query,
            params,
            {
                "월": effective_received_at,
                "주": CSCase.week,
                "식별코드": CSCase.transfer_no,
                "접수일": effective_received_at,
                "회수일": CSCase.return_at,
                "병원명": CSCase.hospital_name,
                "유형(대)": CSCase.category_main,
                "유형(소)": CSCase.category_sub,
                "교품": CSCase.exchange_type,
                "고장 S/N": CSCase.serial_no,
                "LOT": CSCase.lot,
                "Size": CSCase.size,
                "교체 S/N": CSCase.replacement_serial_no,
            },
            {"유형(대)", "유형(소)", "교품"},
        )
        transfer_no = (params.get("transfer_no") or "").strip()
        if transfer_no:
            query = query.filter(CSCase.transfer_no == transfer_no)
        sort_column = {
            "received_at": effective_received_at,
            "week": CSCase.week,
            "transfer_no": CSCase.transfer_no,
        }.get(params.get("sort"), effective_received_at)
        order_fn = desc if params.get("order") == "desc" else asc
        return query.order_by(order_fn(sort_column)).all()

    if section == "quality":
        query = (
            db.query(CSCase, QualityInspection, CXAnalysis)
            .outerjoin(ReturnTransferItem, CSCase.id == ReturnTransferItem.cs_id)
            .outerjoin(
                ReturnTransferRequest,
                ReturnTransferItem.request_id == ReturnTransferRequest.id,
            )
            .outerjoin(
                QualityInspection, CSCase.transfer_no == QualityInspection.transfer_no
            )
            .outerjoin(CXAnalysis, CSCase.transfer_no == CXAnalysis.transfer_no)
            .filter(
                func.lower(func.trim(CSCase.exchange_type)).in_(["o", "cradle"]),
                CSCase.return_at.isnot(None),
                CSCase.product == product,
                or_(
                    ReturnTransferRequest.status == "approved",
                    QualityInspection.id.isnot(None),
                ),
                or_(CSCase.serial_no.is_(None), ~CSCase.serial_no.like("F23%")),
                or_(CSCase.serial_no.is_(None), ~CSCase.serial_no.like("F24%")),
            )
        )
        transfer_no = (params.get("transfer_no") or "").strip()
        if transfer_no:
            query = query.filter(CSCase.transfer_no.contains(transfer_no))
        result, seen = [], set()
        for cs, quality, cx in query.order_by(CSCase.return_at.asc()).all():
            if cs.id not in seen:
                seen.add(cs.id)
                result.append((cs, quality, cx))
        return result

    if section == "lab":
        query = (
            db.query(CSCase, QualityInspection, LabAnalysis)
            .join(
                QualityInspection, CSCase.transfer_no == QualityInspection.transfer_no
            )
            .join(LabAnalysis, CSCase.transfer_no == LabAnalysis.transfer_no)
            .filter(CSCase.product == product)
        )
        transfer_no = (params.get("transfer_no") or "").strip()
        if transfer_no:
            query = query.filter(CSCase.transfer_no == transfer_no)
        return query.order_by(
            QualityInspection.lab_transfer_date.asc(), LabAnalysis.id.asc()
        ).all()

    query = (
        db.query(CSCase, QualityInspection, CXAnalysis, LabAnalysis)
        .outerjoin(
            QualityInspection, CSCase.transfer_no == QualityInspection.transfer_no
        )
        .outerjoin(CXAnalysis, CSCase.transfer_no == CXAnalysis.transfer_no)
        .outerjoin(LabAnalysis, CSCase.transfer_no == LabAnalysis.transfer_no)
        .filter(CSCase.product == product)
    )
    transfer_no = (params.get("transfer_no") or "").strip()
    if transfer_no:
        query = query.filter(CSCase.transfer_no == transfer_no)
    return query.order_by(CSCase.id.desc()).all()


QUALITY_QUERY_FILTER_HEADERS = {
    "transfer_no": "식별코드",
    "serial_no": "S/N",
    "lot": "S/N",
    "rev": "Rev",
    **{field: label for label, field in QUALITY_HEADER_FIELDS.items()},
    "lab_transfer_date": "품질 분석 완료일",
}

EXACT_DROPDOWN_FILTER_HEADERS = {
    "CX",
    "RnD",
    "RnD Team",
    "접수 유형(대)",
    "접수 유형(소)",
    "교품 여부",
    "데이터 손실",
    "불량원인 1",
    "Depth 1-1",
    "Depth 1-2",
    "불량원인 2",
    "Depth 2-1",
    "Depth 2-2",
}


def apply_excel_filters(headers, data, comments, params, section):
    filters = {
        key[3:]: str(value).strip()
        for key, value in params.multi_items()
        if key.startswith("cf_") and str(value).strip()
    }
    if section == "quality":
        for param, header in QUALITY_QUERY_FILTER_HEADERS.items():
            value = str(params.get(param) or "").strip()
            if value:
                filters.setdefault(header, value)
    if not filters:
        return data, comments

    header_indexes = {header: index for index, header in enumerate(headers)}
    quality_cx_headers = (
        "유형(대)",
        "유형(소)",
        "문의사항",
        "CX 원인1",
        "CX Depth1-1",
        "CX Depth1-2",
        "CX 원인2",
        "CX Depth2-1",
        "CX Depth2-2",
    )

    def searchable_value(row, header):
        if section == "quality" and header == "CX 정보":
            return " ".join(
                str(row[header_indexes[name]] or "")
                for name in quality_cx_headers
                if name in header_indexes
            )
        index = header_indexes.get(header)
        return "" if index is None else str(row[index] or "")

    def matches(row):
        for header, expected in filters.items():
            if header in {"흐름", ""}:
                continue
            actual = searchable_value(row, header).strip()
            normalized = expected.lower()
            if normalized in {"null", BLANK_FILTER_VALUE}:
                if actual:
                    return False
            elif normalized == "not_null":
                if not actual:
                    return False
            elif header in EXACT_DROPDOWN_FILTER_HEADERS:
                if actual.strip().lower() != normalized:
                    return False
            elif normalized not in actual.lower():
                return False
        return True

    filtered_data = []
    excel_row_map = {}
    for original_excel_row, row in enumerate(data, start=2):
        if matches(row):
            excel_row_map[original_excel_row] = len(filtered_data) + 2
            filtered_data.append(row)
    filtered_comments = {
        (excel_row_map[row], column): memo
        for (row, column), memo in comments.items()
        if row in excel_row_map
    }
    return filtered_data, filtered_comments


def excel_data(db, request, product, section):
    bp_profile = is_bp_product(product)
    headers = list(SECTION_HEADERS[section])
    if section in {"exchange", "quality"} and bp_profile:
        headers.remove("병원명")
    rows = page_rows(db, request, product, section)
    data, comments = [], {}
    if section == "cx":
        for cs, cx in rows:
            is_cx = (cs.cx_transfer or (cx.cx_yn if cx else "") or "").upper() == "O"
            received_at = cs.received_at or (cx.received_at if cx else None)
            processed_at = cs.processed_at or (cx.processed_at if cx else None)
            serial = cs.serial_no or ""
            data.append(
                [
                    cs.date_only.month if cs.date_only else "",
                    calendar_week(cs.date_only) or "",
                    fmt_date(cs.date_only),
                    cs.hospital_name or "",
                    cs.rev or rev_from_serial(serial) or "",
                    display_transfer_no(cs.transfer_no),
                    cs.cx_transfer or (cx.cx_yn if cx else ""),
                    cs.rnd_transfer or (cx.rnd_yn if cx else ""),
                    cs.rnd_team or (cx.rnd_team if cx else ""),
                    cs.category_main or "",
                    cs.category_sub or "",
                    cs.description or "",
                    cs.exchange_type or "",
                    serial,
                    cs.lot or (serial.split("-")[0] if serial else ""),
                    cs.size or "",
                    fmt_compact(received_at) if is_cx else "-",
                    fmt_compact(processed_at) if is_cx else "-",
                    handling_minutes(received_at, processed_at) if is_cx else "-",
                    cx.result_summary if cx else "",
                    (
                        "기기 미고장 손실"
                        if cx and cx.data_loss_type in {"general", "non_device_failure"}
                        else (
                            "기기 고장 손실"
                            if cx and cx.data_loss_type == "device_failure"
                            else ""
                        )
                    ),
                    cx.data_loss_count if cx and cx.data_loss_type else "",
                    cx.issue1 if cx else "",
                    cx.depth1_1 if cx else "",
                    cx.depth1_2 if cx else "",
                    cx.issue2 if cx else "",
                    cx.depth2_1 if cx else "",
                    cx.depth2_2 if cx else "",
                ]
            )
    elif section == "exchange":
        for cs, _cx in rows:
            received_at = cs.date_only or cs.received_at
            data.append(
                [
                    received_at.month if received_at else "-",
                    f"W{cs.week}" if cs.week else "-",
                    cs.transfer_no or "",
                    fmt_date(received_at),
                    fmt_date(cs.return_at),
                    *(
                        []
                        if is_bp_product(product)
                        else [cs.hospital_name or ""]
                    ),
                    cs.category_main or "",
                    cs.category_sub or "",
                    cs.exchange_type or "",
                    cs.serial_no or "",
                    cs.lot or "",
                    cs.size or "",
                    lead_days(received_at, cs.return_at),
                    cs.replacement_serial_no or "",
                ]
            )
    elif section == "quality":
        quality_ids = [quality.id for _cs, quality, _cx in rows if quality]
        memo_map = (
            {
                (memo.quality_id, memo.field_name): memo.memo
                for memo in db.query(QualityCellMemo)
                .filter(QualityCellMemo.quality_id.in_(quality_ids))
                .all()
            }
            if quality_ids
            else {}
        )
        for excel_row, (cs, quality, _cx) in enumerate(rows, start=2):
            values = [
                fmt_date(cs.received_at or cs.date_only),
                fmt_date(cs.return_at),
                fmt_date(quality.return_transfer_date) if quality else "",
                *([cs.hospital_name or ""] if not bp_profile else []),
                cs.transfer_no or "",
                cs.serial_no or "",
                resolved_quality_rev(cs, quality),
                cs.rev or "",
                cs.category_main or "",
                cs.category_sub or "",
                cs.description or "",
                _cx.issue1 if _cx else "",
                _cx.depth1_1 if _cx else "",
                _cx.depth1_2 if _cx else "",
                _cx.issue2 if _cx else "",
                _cx.depth2_1 if _cx else "",
                _cx.depth2_2 if _cx else "",
            ]
            values.extend(
                (
                    (getattr(quality, field) or "대기")
                    if quality and field == "quality_result"
                    else getattr(quality, field) if quality else ""
                )
                for field in QUALITY_FIELDS
            )
            values.append(fmt_date(quality.lab_transfer_date) if quality else "")
            data.append(values)
            if quality:
                memo_start_column = 18 if not bp_profile else 17
                for column, field in enumerate(QUALITY_FIELDS, start=memo_start_column):
                    memo = memo_map.get((quality.id, field))
                    if memo:
                        comments[(excel_row, column)] = memo
    elif section == "lab":
        for cs, quality, lab in rows:
            try:
                attachments = json.loads(lab.attachment or "[]")
            except (TypeError, ValueError):
                attachments = [lab.attachment] if lab.attachment else []
            data.append(
                [
                    cs.transfer_no or "",
                    lab.serial_no or cs.serial_no or "",
                    fmt_date(quality.lab_transfer_date),
                    lab.root_cause or "",
                    lab.analysis_result or "",
                    lab.defect_type or "",
                    lab.remark or "",
                    fmt_date(lab.analysis_completed_date),
                    f"{len(attachments)}개",
                ]
            )
    else:
        for cs, quality, cx, lab in rows:
            data.append(
                [
                    cs.transfer_no or "",
                    fmt_date(cs.received_at),
                    cs.hospital_name or "",
                    cs.category_main or "",
                    cs.category_sub or "",
                    cs.description or "",
                    cs.serial_no or "",
                    cs.lot or "",
                    cs.size or "",
                    cs.exchange_type or "X",
                    cx.result_summary if cx else "",
                    quality_status(cs, quality),
                    lab_status(quality, lab),
                    lead_days(cs.received_at, cs.return_at),
                    lead_days(
                        cs.return_at, quality.lab_transfer_date if quality else None
                    ),
                    lead_days(
                        quality.lab_transfer_date if quality else None,
                        lab.analysis_completed_date if lab else None,
                    ),
                    quality.quality_result if quality else "",
                    lab.analysis_result if lab else "",
                    fmt_date(lab.analysis_completed_date) if lab else "",
                ]
            )
    if section == "cx" and bp_profile:
        excluded_headers = {"병원명", "데이터 손실", "손실 건수"}
        kept_indexes = [
            index
            for index, header in enumerate(headers)
            if header not in excluded_headers
        ]
        headers = [headers[index] for index in kept_indexes]
        data = [[row[index] for index in kept_indexes] for row in data]

    data, comments = apply_excel_filters(
        headers, data, comments, request.query_params, section
    )
    return headers, data, comments


def style_workbook(worksheet):
    fill = PatternFill("solid", fgColor="E8F1F5")
    for cell in worksheet[1]:
        cell.font = Font(bold=True, color="24364B")
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for index, column in enumerate(worksheet.columns, start=1):
        longest = min(
            max((len(str(cell.value or "")) for cell in column), default=0) + 2, 45
        )
        worksheet.column_dimensions[get_column_letter(index)].width = max(10, longest)


def build_download_workbook(title, headers, rows, comments):
    """Build a styled workbook in one forward-only pass.

    Download data is already materialized for filtering, so write-only mode
    avoids constructing a second full worksheet object and then scanning every
    cell twice for alignment and width calculation.
    """
    workbook = Workbook(write_only=True)
    worksheet = workbook.create_sheet(title=title)
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = (
        f"A1:{get_column_letter(len(headers))}{max(len(rows) + 1, 1)}"
    )

    widths = [min(max(len(str(header or "")) + 2, 10), 45) for header in headers]
    for row in rows:
        for index, value in enumerate(row):
            if index < len(widths):
                widths[index] = min(max(widths[index], len(str(value or "")) + 2), 45)
    for index, width in enumerate(widths, start=1):
        worksheet.column_dimensions[get_column_letter(index)].width = width

    header_fill = PatternFill("solid", fgColor="E8F1F5")
    header_cells = []
    for value in headers:
        cell = WriteOnlyCell(worksheet, value=value)
        cell.font = Font(bold=True, color="24364B")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        header_cells.append(cell)
    worksheet.append(header_cells)

    body_alignment = Alignment(vertical="top", wrap_text=True)
    for excel_row, values in enumerate(rows, start=2):
        cells = []
        for column, value in enumerate(values, start=1):
            cell = WriteOnlyCell(worksheet, value=value)
            cell.alignment = body_alignment
            memo = comments.get((excel_row, column))
            if memo:
                cell.comment = Comment(memo, "Service Monitor")
            cells.append(cell)
        worksheet.append(cells)
    return workbook


@router.get("/{product}/{section}/excel-download")
def excel_download(request: Request, product: str, section: str):
    if section not in SECTIONS:
        raise HTTPException(404, "Excel 다운로드를 지원하지 않는 페이지입니다.")
    if section == "result":
        from app.routers.result import download_master_excel

        return download_master_excel(request, product)
    db = SessionLocal()
    try:
        headers, rows, comments = excel_data(db, request, product, section)
        title = {
            "cx": "CX 분석",
            "exchange": "교품 회수",
            "quality": "품질 분석",
            "lab": "연구소 분석",
            "result": "최종 결과",
        }[section]
        workbook = build_download_workbook(title, headers, rows, comments)
        output = BytesIO()
        workbook.save(output)
        output.seek(0)
        return StreamingResponse(
            output,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={
                "Content-Disposition": dated_excel_disposition(
                    product,
                    {
                        "cx": "CX 분석",
                        "exchange": "교품／회수 리스트",
                        "quality": "회수품 검토 이력",
                        "lab": "연구소 분석",
                    }[section],
                )
            },
        )
    finally:
        db.close()


def set_string_fields(target, row, mapping):
    for header, field in mapping.items():
        if header in row:
            setattr(target, field, clean_text(row[header]))


def get_or_create_analysis(db, code):
    item = db.query(CXAnalysis).filter(CXAnalysis.transfer_no == code).first()
    if not item:
        item = CXAnalysis(transfer_no=code)
        db.add(item)
    return item


def get_or_create_quality(db, code):
    item = (
        db.query(QualityInspection)
        .filter(QualityInspection.transfer_no == code)
        .first()
    )
    if not item:
        item = QualityInspection(transfer_no=code)
        db.add(item)
        db.flush()
    return item


def import_cx_row(db, product, row, default_year=None):
    code = clean_text(row.get("식별코드"))
    size = clean_text(row.get("Size")) if "Size" in row else None
    if size and not re.fullmatch(r"\d{1,2}", size):
        raise ValueError("Size는 1~2자리 숫자만 입력할 수 있습니다.")
    for header in ("CX", "RnD"):
        value = (clean_text(row.get(header)) or "").upper() if header in row else ""
        if value not in {"", "O"}:
            raise ValueError(f"{header}는 O 또는 빈칸만 입력할 수 있습니다.")
    if "RnD Team" in row:
        team = clean_text(row.get("RnD Team"))
        team = {
            "APP": "App",
            "Admin개선": "Admin 개선",
        }.get(team, team)
        if team and team not in {"App", "Server", "F/W", "H/W", "VSP", "Admin 개선"}:
            raise ValueError("RnD Team 값을 확인해 주세요.")
        row["RnD Team"] = team
    if "교품 여부" in row:
        exchange_type = clean_text(row.get("교품 여부"))
        exchange_type = {
            "o": "O",
            "x": "X",
            "cradle": "Cradle",
        }.get(str(exchange_type or "").lower(), exchange_type)
        if exchange_type and exchange_type not in {"O", "X", "Cradle"}:
            raise ValueError("교품 여부는 O, X, Cradle 또는 빈칸만 입력할 수 있습니다.")
        row["교품 여부"] = exchange_type
    loss_type = clean_text(row.get("데이터 손실")) if "데이터 손실" in row else None
    loss_type_map = {
        "기기 미고장 손실": "non_device_failure",
        "기기 고장 손실": "device_failure",
        "일반 손실": "non_device_failure",
        "general": "non_device_failure",
        "non_device_failure": "non_device_failure",
        "device_failure": "device_failure",
    }
    if loss_type and loss_type not in loss_type_map:
        raise ValueError(
            "데이터 손실은 기기 고장 손실, 기기 미고장 손실 또는 빈칸만 입력할 수 있습니다."
        )
    loss_count = clean_text(row.get("손실 건수")) if "손실 건수" in row else None
    if loss_count:
        try:
            loss_count = int(loss_count)
        except (TypeError, ValueError):
            raise ValueError("손실 건수는 1~999 사이의 숫자여야 합니다.")
        if not 1 <= loss_count <= 999:
            raise ValueError("손실 건수는 1~999 사이의 숫자여야 합니다.")
        if not loss_type:
            raise ValueError("손실 건수를 입력하려면 데이터 손실 분류가 필요합니다.")
    case = (
        db.query(CSCase)
        .filter(CSCase.product == product, CSCase.transfer_no == code)
        .first()
    )
    if not case:
        case = CSCase(product=product, transfer_no=code)
        db.add(case)
    set_string_fields(
        case,
        row,
        {
            "병원명": "hospital_name",
            "REV 구분": "rev",
            "접수 유형(대)": "category_main",
            "접수 유형(소)": "category_sub",
            "문의사항": "description",
            "교품 여부": "exchange_type",
            "고장 S/N": "serial_no",
            "S/N": "serial_no",
            "LOT": "lot",
            "Size": "size",
        },
    )
    if "일자" in row:
        case.date_only = parse_datetime(row["일자"], default_year=default_year)
        case.week = calendar_week(case.date_only)
    if "S/N" in row:
        serial = (case.serial_no or "").upper()
        resolved_rev = rev_from_serial(serial)
        if resolved_rev:
            case.rev = resolved_rev
        elif "REV 구분" not in row:
            case.rev = None
    cx_headers = {
        "CX",
        "RnD",
        "RnD Team",
        "접수",
        "처리",
        "처리 결과",
        "데이터 손실",
        "손실 건수",
        "불량원인 1",
        "Depth 1-1",
        "Depth 1-2",
        "불량원인 2",
        "Depth 2-1",
        "Depth 2-2",
    }
    analysis = db.query(CXAnalysis).filter(CXAnalysis.transfer_no == code).first()
    has_cx_value = any(
        clean_text(row.get(header)) for header in cx_headers if header in row
    )
    if analysis or has_cx_value:
        analysis = analysis or get_or_create_analysis(db, code)
        set_string_fields(
            analysis,
            row,
            {
                "CX": "cx_yn",
                "RnD": "rnd_yn",
                "RnD Team": "rnd_team",
                "처리 결과": "result_summary",
                "불량원인 1": "issue1",
                "Depth 1-1": "depth1_1",
                "Depth 1-2": "depth1_2",
                "불량원인 2": "issue2",
                "Depth 2-1": "depth2_1",
                "Depth 2-2": "depth2_2",
            },
        )
        is_cx = (analysis.cx_yn or "").upper() == "O"
        if is_cx and "접수" in row:
            analysis.received_at = parse_datetime(
                row["접수"], compact=True, default_year=default_year
            )
            case.received_at = analysis.received_at
        if is_cx and "처리" in row:
            analysis.processed_at = parse_datetime(
                row["처리"], compact=True, default_year=default_year
            )
            case.processed_at = analysis.processed_at
        if "CX" in row:
            case.cx_transfer = analysis.cx_yn
        if "RnD" in row:
            case.rnd_transfer = analysis.rnd_yn
        if "RnD Team" in row:
            case.rnd_team = analysis.rnd_team
        if "데이터 손실" in row:
            if loss_type and not case.date_only:
                raise ValueError("데이터 손실을 지정하려면 일자가 필요합니다.")
            analysis.data_loss_type = loss_type_map.get(loss_type)
            analysis.data_loss_count = (loss_count or 1) if loss_type else None
        elif "손실 건수" in row and loss_count:
            analysis.data_loss_count = loss_count


def import_exchange_row(db, product, row):
    code = clean_text(row.get("식별코드"))
    case = (
        db.query(CSCase)
        .filter(CSCase.product == product, CSCase.transfer_no == code)
        .first()
    )
    if not case:
        raise ValueError("기존 CX 접수를 찾을 수 없습니다.")
    set_string_fields(
        case,
        row,
        {
            "병원명": "hospital_name",
            "유형(대)": "category_main",
            "유형(소)": "category_sub",
            "교품": "exchange_type",
            "S/N": "serial_no",
            "송장 번호": "invoice_no",
            "교체 S/N": "replacement_serial_no",
            "대치품 S/N": "replacement_serial_no",
            "LOT": "lot",
            "Size": "size",
        },
    )
    if "접수일" in row:
        case.received_at = parse_datetime(row["접수일"])
        case.week = calendar_week(case.received_at)
    if "회수일" in row:
        case.return_at = parse_datetime(row["회수일"])
    if "교체 S/N" in row or "대치품 S/N" in row:
        case.replacement_serial_no_manual = True


def import_quality_row(db, product, row, cells):
    code = clean_text(row.get("식별코드"))
    case = (
        db.query(CSCase)
        .filter(CSCase.product == product, CSCase.transfer_no == code)
        .first()
    )
    if not case:
        raise ValueError("기존 CX 접수를 찾을 수 없습니다.")
    quality = get_or_create_quality(db, code)
    if "공정 Rev" in row:
        quality.rev = clean_text(row["공정 Rev"])
    elif "Rev" in row:
        quality.rev = clean_text(row["Rev"])
    if "HW Rev" in row:
        case.rev = clean_text(row["HW Rev"])
    if "회수품 이관일" in row:
        parsed = parse_datetime(row["회수품 이관일"])
        quality.return_transfer_date = parsed.date() if parsed else None
    for header, field in QUALITY_HEADER_FIELDS.items():
        if header not in row:
            continue
        value = clean_text(row[header]) if isinstance(row[header], str) else row[header]
        setattr(quality, field, value)
        comment = cells[header].comment.text.strip() if cells[header].comment else ""
        if comment:
            memo = (
                db.query(QualityCellMemo)
                .filter(
                    QualityCellMemo.quality_id == quality.id,
                    QualityCellMemo.field_name == field,
                )
                .first()
            )
            if not memo:
                memo = QualityCellMemo(quality_id=quality.id, field_name=field)
                db.add(memo)
            memo.memo = comment
    if "품질 분석 완료일" in row or "연구소 이관일" in row:
        parsed = parse_datetime(row.get("품질 분석 완료일") or row.get("연구소 이관일"))
        quality.lab_transfer_date = parsed.date() if parsed else None


def import_lab_row(db, product, row):
    code = clean_text(row.get("식별코드"))
    case = (
        db.query(CSCase)
        .filter(CSCase.product == product, CSCase.transfer_no == code)
        .first()
    )
    if not case:
        raise ValueError("기존 CX 접수를 찾을 수 없습니다.")
    quality = (
        db.query(QualityInspection)
        .filter(QualityInspection.transfer_no == code)
        .first()
    )
    if not quality:
        raise ValueError("품질 분석 데이터가 없습니다.")
    lab = db.query(LabAnalysis).filter(LabAnalysis.transfer_no == code).first()
    if not lab:
        lab = LabAnalysis(transfer_no=code, serial_no=case.serial_no)
        db.add(lab)
    set_string_fields(
        lab,
        row,
        {
            "S/N": "serial_no",
            "불량 원인": "root_cause",
            "분석 결과": "analysis_result",
            "불량 유형": "defect_type",
            "비고": "remark",
        },
    )
    if "품질 분석 완료일" in row or "연구소 이관일" in row:
        parsed = parse_datetime(row.get("품질 분석 완료일") or row.get("연구소 이관일"))
        quality.lab_transfer_date = parsed.date() if parsed else None
    if "분석 완료일" in row:
        parsed = parse_datetime(row["분석 완료일"])
        lab.analysis_completed_date = parsed.date() if parsed else None


@router.post("/{product}/{section}/excel-upload")
async def excel_upload(product: str, section: str, file: UploadFile = File(...)):
    if section not in SECTIONS or section == "result":
        raise HTTPException(400, "이 페이지는 Excel 업로드를 지원하지 않습니다.")
    if not (file.filename or "").lower().endswith((".xlsx", ".xlsm")):
        raise HTTPException(400, "Excel 파일(.xlsx, .xlsm)만 업로드할 수 있습니다.")
    contents = await file.read()
    if not contents:
        raise HTTPException(400, "빈 파일은 업로드할 수 없습니다.")
    if len(contents) > MAX_UPLOAD_BYTES:
        raise HTTPException(400, "Excel 파일은 10MB 이하만 업로드할 수 있습니다.")
    try:
        workbook = load_workbook(BytesIO(contents), data_only=True)
    except Exception as error:
        raise HTTPException(400, "Excel 파일을 읽을 수 없습니다.") from error
    worksheet = workbook.active
    headers = [str(cell.value or "").strip() for cell in worksheet[1]]
    if "식별코드" not in headers:
        raise HTTPException(400, "첫 번째 행에 '식별코드' 열이 필요합니다.")
    if set(headers).intersection(SECTION_HEADERS[section]) == {"식별코드"}:
        raise HTTPException(400, "현재 페이지에서 사용하는 수정 가능 열이 없습니다.")
    import_year = None
    if section == "cx":
        date_indexes = [
            index
            for index, header in enumerate(headers)
            if header in {"일자", "접수", "처리"}
        ]
        years = []
        for cells in worksheet.iter_rows(min_row=2, max_col=len(headers)):
            for index in date_indexes:
                value = cells[index].value
                if isinstance(value, (datetime, date)):
                    years.append(value.year)
                elif value:
                    match = re.search(r"(?<!\d)(20\d{2})(?!\d)", str(value))
                    if match:
                        years.append(int(match.group(1)))
        if years:
            import_year = Counter(years).most_common(1)[0][0]
    db = SessionLocal()
    updated, skipped, errors = 0, 0, []
    try:
        for excel_row, cells in enumerate(
            worksheet.iter_rows(min_row=2, max_col=len(headers)), start=2
        ):
            row = {header: cell.value for header, cell in zip(headers, cells) if header}
            cell_map = {header: cell for header, cell in zip(headers, cells) if header}
            code = clean_text(row.get("식별코드"))
            if not code and not any(clean_text(value) for value in row.values()):
                skipped += 1
                continue
            if not code or str(code).strip().upper() == "N/A":
                row["식별코드"] = (
                    f"__NA__{import_year or datetime.now().year}_{excel_row:06d}"
                )
            try:
                with db.begin_nested():
                    if section == "cx":
                        import_cx_row(db, product, row, default_year=import_year)
                    elif section == "exchange":
                        import_exchange_row(db, product, row)
                    elif section == "quality":
                        import_quality_row(db, product, row, cell_map)
                    else:
                        import_lab_row(db, product, row)
                    db.flush()
                updated += 1
            except (TypeError, ValueError) as error:
                errors.append(f"{excel_row}행({code}): {error}")
        db.commit()
        return {
            "status": "ok",
            "updated": updated,
            "skipped": skipped,
            "errors": errors[:20],
        }
    finally:
        db.close()