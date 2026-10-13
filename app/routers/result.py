from fastapi import APIRouter, Request
from fastapi.templating import Jinja2Templates
from fastapi.responses import HTMLResponse, JSONResponse, StreamingResponse
from sqlalchemy.orm import Session
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from io import BytesIO
from app.database import SessionLocal
from app.models.cs import CSCase
from app.models.cx import CXAnalysis
from app.models.quality import QualityInspection
from app.models.lab import LabAnalysis
from app.models.quality_memo import QualityCellMemo
from app.services.quality_schema import QUALITY_COLUMNS, QUALITY_LABELS
from app.services.cx_categories import rev_from_serial
from app.utils.query_filters import apply_column_filters
from app.utils.downloads import dated_excel_disposition
from sqlalchemy import String, and_, asc, cast, desc, func, or_

router = APIRouter()
templates = Jinja2Templates(directory="app/templates")
EXCHANGE_TYPES = ("O", "Cradle", "cradle")


def _case_state(cs, quality, cx, lab):
    exchange_required = cs.exchange_type in EXCHANGE_TYPES
    cx_required = (cs.cx_transfer or "").upper() == "O" or bool(
        cx and (cx.cx_yn or "").upper() == "O"
    )
    cx_done = cx_required and bool((cx.processed_at if cx else None) or cs.processed_at)
    recovery_done = bool(cs.return_at)
    quality_started = exchange_required and recovery_done and bool(quality)
    quality_done = bool(
        quality_started
        and quality.quality_result in {"OK", "NG"}
        and quality.lab_transfer_date
    )
    rnd_required = quality_done and bool(quality and quality.lab_transfer_date)
    rnd_started = rnd_required and bool(lab)
    rnd_done = bool(rnd_started and lab.analysis_completed_date)
    if cx_required and not cx_done:
        current_stage = "CX 처리"
    elif not exchange_required:
        current_stage = "CX 완료" if cx_done else "접수 완료"
    elif rnd_required:
        current_stage = "완료" if rnd_done else "R&D 분석"
    elif not recovery_done:
        current_stage = "회수 대기"
    elif quality_started:
        current_stage = "품질 완료" if quality_done else "품질 분석"
    elif recovery_done:
        current_stage = "품질 대기"
    elif cx_done:
        current_stage = "CX 완료"
    else:
        current_stage = "CX 처리"
    checkpoints = (
        ([True, cx_done] if cx_required else [True])
        if not exchange_required
        else [
            True,
            *([cx_done] if cx_required else []),
            recovery_done,
            quality_done,
            (rnd_done if rnd_required else quality_done),
        ]
    )
    progress = round(sum(checkpoints) / len(checkpoints) * 100)
    return {
        "current_stage": current_stage,
        "progress": progress,
        "cx_status": (
            "해당 없음" if not cx_required else "완료" if cx_done else "처리 중"
        ),
        "exchange_status": (
            "해당 없음"
            if not exchange_required
            else "회수 완료" if recovery_done else "회수 대기"
        ),
        "quality_status": (
            "해당 없음"
            if not exchange_required
            else "완료" if quality_done else "분석 중" if quality_started else "대기"
        ),
        "rnd_status": (
            "해당 없음"
            if not exchange_required or not rnd_required
            else "완료" if rnd_done else "분석 중" if rnd_started else "대기"
        ),
        "exchange_required": exchange_required,
        "cx_required": cx_required,
    }


def _hw_rev(cs):
    return cs.rev or rev_from_serial(cs.serial_no) or ""


def _display(value):
    if value is None:
        return ""
    if hasattr(value, "isoformat"):
        return (
            value.isoformat(sep=" ", timespec="minutes")
            if hasattr(value, "hour")
            else value.isoformat()
        )
    return value


def _model_values(model, excluded=("id",)):
    if not model:
        return {}
    return {
        column.name: _display(getattr(model, column.name))
        for column in model.__table__.columns
        if column.name not in excluded
    }


@router.get("/{product}/result")
def result_list(request: Request, product: str):
    db: Session = SessionLocal()
    keyword = request.query_params.get("keyword", "").strip()
    stage = request.query_params.get("stage", "all").strip().lower()
    order = request.query_params.get("order", "desc")
    order = order if order in {"asc", "desc"} else "desc"
    if stage not in {"all", "cx", "recovery", "quality", "rnd", "complete"}:
        stage = "all"
    try:
        page = max(1, int(request.query_params.get("page", 1)))
    except (TypeError, ValueError):
        page = 1
    try:
        requested_limit = int(request.query_params.get("limit", 50))
    except (TypeError, ValueError):
        requested_limit = 50
    limit = requested_limit if requested_limit in {20, 50, 100} else 50

    offset = (page - 1) * limit

    query = (
        db.query(CSCase, QualityInspection, CXAnalysis, LabAnalysis)
        .outerjoin(
            QualityInspection, CSCase.transfer_no == QualityInspection.transfer_no
        )
        .outerjoin(CXAnalysis, CSCase.transfer_no == CXAnalysis.transfer_no)
        .outerjoin(LabAnalysis, CSCase.transfer_no == LabAnalysis.transfer_no)
        .filter(CSCase.product == product)
    )
    exchange_filter = CSCase.exchange_type.in_(EXCHANGE_TYPES)
    non_exchange_filter = or_(
        CSCase.exchange_type.is_(None), ~CSCase.exchange_type.in_(EXCHANGE_TYPES)
    )
    cx_required_filter = or_(
        func.upper(CSCase.cx_transfer) == "O", func.upper(CXAnalysis.cx_yn) == "O"
    )
    cx_complete_filter = or_(
        CXAnalysis.processed_at.isnot(None), CSCase.processed_at.isnot(None)
    )

    if stage == "cx":
        query = query.filter(cx_required_filter, ~cx_complete_filter)
    elif stage == "recovery":
        query = query.filter(exchange_filter, CSCase.return_at.is_(None))
    elif stage == "quality":
        query = query.filter(
            exchange_filter,
            CSCase.return_at.isnot(None),
            or_(
                QualityInspection.id.is_(None),
                QualityInspection.quality_result.is_(None),
            ),
        )
    elif stage == "rnd":
        query = query.filter(
            exchange_filter,
            QualityInspection.lab_transfer_date.isnot(None),
            or_(
                LabAnalysis.id.is_(None), LabAnalysis.analysis_completed_date.is_(None)
            ),
        )
    elif stage == "complete":
        query = query.filter(
            or_(
                and_(non_exchange_filter, cx_complete_filter),
                and_(
                    exchange_filter,
                    QualityInspection.quality_result.isnot(None),
                    or_(
                        QualityInspection.lab_transfer_date.is_(None),
                        LabAnalysis.analysis_completed_date.isnot(None),
                    ),
                ),
            )
        )

    if keyword:
        token = f"%{keyword}%"
        search_columns = [
            *CSCase.__table__.columns,
            *CXAnalysis.__table__.columns,
            *QualityInspection.__table__.columns,
            *LabAnalysis.__table__.columns,
        ]
        query = query.filter(
            or_(*[cast(column, String).ilike(token) for column in search_columns])
        )
    query = apply_column_filters(
        query,
        request.query_params,
        {
            "식별코드": CSCase.transfer_no,
            "고객·기기": [CSCase.hospital_name, CSCase.serial_no, CSCase.rev],
            "접수 유형": [CSCase.category_main, CSCase.category_sub],
            "CX": [CSCase.cx_transfer, CXAnalysis.cx_yn],
            "교품·회수": [CSCase.exchange_type, CSCase.return_yn, CSCase.return_at],
            "품질": [
                QualityInspection.quality_result,
                QualityInspection.lab_transfer_date,
            ],
            "R&D": [LabAnalysis.analysis_result, LabAnalysis.analysis_completed_date],
        },
    )

    total_count = query.count()

    total_pages = (total_count + limit - 1) // limit

    order_fn = asc if order == "asc" else desc
    rows = (
        query.order_by(
            CSCase.date_only.is_(None), order_fn(CSCase.date_only), order_fn(CSCase.id)
        )
        .offset(offset)
        .limit(limit)
        .all()
    )

    result_rows = []

    for cs, quality, cx, lab in rows:
        if cs.return_at and not quality:
            quality_status = "대기 중"
        elif quality and not quality.lab_transfer_date:
            quality_status = "분석 중"
        elif quality and quality.lab_transfer_date:
            quality_status = "분석 완료"
        else:
            quality_status = "-"
        if quality and quality.lab_transfer_date and not lab:
            rnd_status = "대기 중"
        elif lab and not lab.analysis_completed_date:
            rnd_status = "분석 중"
        elif lab and lab.analysis_completed_date:
            rnd_status = "분석 완료"
        else:
            rnd_status = "-"

        lead_receive_to_return = (
            (cs.return_at.date() - cs.received_at.date()).days
            if cs.received_at and cs.return_at
            else None
        )

        lead_return_to_quality = (
            (quality.lab_transfer_date - cs.return_at.date()).days
            if quality and quality.lab_transfer_date and cs.return_at
            else None
        )

        lead_quality_to_rnd = (
            (lab.analysis_completed_date - quality.lab_transfer_date).days
            if lab
            and lab.analysis_completed_date
            and quality
            and quality.lab_transfer_date
            else None
        )

        state = _case_state(cs, quality, cx, lab)
        result_rows.append(
            {
                "cs": cs,
                "quality": quality,
                "cx": cx,
                "lab": lab,
                "quality_status": quality_status,
                "rnd_status": rnd_status,
                "lead_receive_to_return": lead_receive_to_return,
                "lead_return_to_quality": lead_return_to_quality,
                "lead_quality_to_rnd": lead_quality_to_rnd,
                "hw_rev": _hw_rev(cs),
                **state,
            }
        )

    product_query = db.query(CSCase).filter(CSCase.product == product)
    master_metrics = {
        "all": product_query.count(),
        "cx": product_query.outerjoin(
            CXAnalysis, CSCase.transfer_no == CXAnalysis.transfer_no
        )
        .filter(
            or_(
                func.upper(CSCase.cx_transfer) == "O",
                func.upper(CXAnalysis.cx_yn) == "O",
            ),
            CXAnalysis.processed_at.is_(None),
            CSCase.processed_at.is_(None),
        )
        .count(),
        "recovery": product_query.filter(
            CSCase.exchange_type.in_(EXCHANGE_TYPES), CSCase.return_at.is_(None)
        ).count(),
        "quality": product_query.outerjoin(
            QualityInspection, CSCase.transfer_no == QualityInspection.transfer_no
        )
        .filter(
            CSCase.exchange_type.in_(EXCHANGE_TYPES),
            CSCase.return_at.isnot(None),
            or_(
                QualityInspection.id.is_(None),
                QualityInspection.quality_result.is_(None),
            ),
        )
        .count(),
        "rnd": product_query.join(
            QualityInspection, CSCase.transfer_no == QualityInspection.transfer_no
        )
        .outerjoin(LabAnalysis, CSCase.transfer_no == LabAnalysis.transfer_no)
        .filter(
            CSCase.exchange_type.in_(EXCHANGE_TYPES),
            QualityInspection.lab_transfer_date.isnot(None),
            or_(
                LabAnalysis.id.is_(None), LabAnalysis.analysis_completed_date.is_(None)
            ),
        )
        .count(),
        "complete": product_query.outerjoin(
            CXAnalysis, CSCase.transfer_no == CXAnalysis.transfer_no
        )
        .outerjoin(
            QualityInspection, CSCase.transfer_no == QualityInspection.transfer_no
        )
        .outerjoin(LabAnalysis, CSCase.transfer_no == LabAnalysis.transfer_no)
        .filter(
            or_(
                and_(
                    or_(
                        CSCase.exchange_type.is_(None),
                        ~CSCase.exchange_type.in_(EXCHANGE_TYPES),
                    ),
                    or_(
                        CXAnalysis.processed_at.isnot(None),
                        CXAnalysis.result_summary.isnot(None),
                    ),
                ),
                and_(
                    CSCase.exchange_type.in_(EXCHANGE_TYPES),
                    QualityInspection.quality_result.isnot(None),
                    or_(
                        QualityInspection.lab_transfer_date.is_(None),
                        LabAnalysis.analysis_completed_date.isnot(None),
                    ),
                ),
            )
        )
        .count(),
    }

    response = templates.TemplateResponse(
        "result_list.html",
        {
            "request": request,
            "rows": result_rows,
            "product": product,
            "page": page,
            "limit": limit,
            "total_pages": total_pages,
            "total_count": total_count,
            "keyword": keyword,
            "stage": stage,
            "master_metrics": master_metrics,
            "order": order,
            "quality_labels": QUALITY_LABELS,
        },
    )

    db.close()

    return response


@router.get("/{product}/api/result/master/{case_id}")
def result_master_detail(product: str, case_id: int):
    db = SessionLocal()
    try:
        record = (
            db.query(CSCase, QualityInspection, CXAnalysis, LabAnalysis)
            .outerjoin(
                QualityInspection, CSCase.transfer_no == QualityInspection.transfer_no
            )
            .outerjoin(CXAnalysis, CSCase.transfer_no == CXAnalysis.transfer_no)
            .outerjoin(LabAnalysis, CSCase.transfer_no == LabAnalysis.transfer_no)
            .filter(CSCase.id == case_id, CSCase.product == product)
            .first()
        )
        if not record:
            return JSONResponse(
                {"detail": "마스터 데이터를 찾을 수 없습니다."}, status_code=404
            )
        cs, quality, cx, lab = record
        state = _case_state(cs, quality, cx, lab)
        quality_values = _model_values(quality)
        quality_values["hw_rev"] = _display(_hw_rev(cs))
        basic_values = _model_values(cs)
        basic_values["rev"] = _display(_hw_rev(cs))
        if quality:
            memos = (
                db.query(QualityCellMemo)
                .filter(QualityCellMemo.quality_id == quality.id)
                .all()
            )
            quality_values["cell_memos"] = "\n".join(
                f"{memo.field_name}: {memo.memo}" for memo in memos if memo.memo
            )
        timeline = [
            {
                "label": "접수",
                "date": _display(cs.received_at),
                "done": bool(cs.received_at),
            },
            {
                "label": "CX 처리",
                "date": (
                    "해당 없음"
                    if not state["cx_required"]
                    else _display((cx.processed_at if cx else None) or cs.processed_at)
                ),
                "done": state["cx_status"] in {"완료", "해당 없음"},
            },
            {
                "label": "회수",
                "date": (
                    "해당 없음"
                    if not state["exchange_required"]
                    else _display(cs.return_at)
                ),
                "done": state["exchange_status"] in {"회수 완료", "해당 없음"},
            },
            {
                "label": "품질 분석",
                "date": (
                    "해당 없음"
                    if not state["exchange_required"]
                    else _display(quality.lab_transfer_date if quality else None)
                ),
                "done": state["quality_status"] in {"완료", "해당 없음"},
            },
            {
                "label": "R&D 분석",
                "date": (
                    "해당 없음"
                    if state["rnd_status"] == "해당 없음"
                    else _display(lab.analysis_completed_date if lab else None)
                ),
                "done": state["rnd_status"] in {"완료", "해당 없음"},
            },
        ]
        return {
            "case_id": cs.id,
            "transfer_no": cs.transfer_no or "",
            "serial_no": cs.serial_no or "",
            "hospital_name": cs.hospital_name or "",
            "state": state,
            "timeline": timeline,
            "basic": basic_values,
            "cx": _model_values(cx),
            "exchange": {
                key: _display(getattr(cs, key))
                for key in [
                    "exchange_type",
                    "return_yn",
                    "return_at",
                    "invoice_no",
                    "replacement_serial_no",
                    "serial_no",
                    "lot",
                    "size",
                ]
            },
            "quality": quality_values,
            "rnd": _model_values(lab),
        }
    finally:
        db.close()


@router.get("/{product}/result/excel")
def download_master_excel(request: Request, product: str):
    keyword = request.query_params.get("keyword", "").strip()
    stage = request.query_params.get("stage", "all").strip().lower()
    db = SessionLocal()
    query = (
        db.query(CSCase, CXAnalysis, QualityInspection, LabAnalysis)
        .outerjoin(CXAnalysis, CSCase.transfer_no == CXAnalysis.transfer_no)
        .outerjoin(
            QualityInspection, CSCase.transfer_no == QualityInspection.transfer_no
        )
        .outerjoin(LabAnalysis, CSCase.transfer_no == LabAnalysis.transfer_no)
        .filter(CSCase.product == product)
    )
    exchange_filter = CSCase.exchange_type.in_(EXCHANGE_TYPES)
    cx_required_filter = or_(
        func.upper(CSCase.cx_transfer) == "O", func.upper(CXAnalysis.cx_yn) == "O"
    )
    cx_complete_filter = or_(
        CXAnalysis.processed_at.isnot(None), CSCase.processed_at.isnot(None)
    )
    if stage == "cx":
        query = query.filter(cx_required_filter, ~cx_complete_filter)
    elif stage == "recovery":
        query = query.filter(exchange_filter, CSCase.return_at.is_(None))
    elif stage == "quality":
        query = query.filter(
            exchange_filter,
            CSCase.return_at.isnot(None),
            or_(
                QualityInspection.id.is_(None),
                QualityInspection.quality_result.is_(None),
            ),
        )
    elif stage == "rnd":
        query = query.filter(
            exchange_filter,
            QualityInspection.lab_transfer_date.isnot(None),
            or_(
                LabAnalysis.id.is_(None), LabAnalysis.analysis_completed_date.is_(None)
            ),
        )
    elif stage == "complete":
        query = query.filter(
            or_(
                and_(
                    or_(CSCase.exchange_type.is_(None), ~exchange_filter),
                    cx_complete_filter,
                ),
                and_(
                    exchange_filter,
                    QualityInspection.quality_result.isnot(None),
                    or_(
                        QualityInspection.lab_transfer_date.is_(None),
                        LabAnalysis.analysis_completed_date.isnot(None),
                    ),
                ),
            )
        )
    if keyword:
        token = f"%{keyword}%"
        columns = [
            *CSCase.__table__.columns,
            *CXAnalysis.__table__.columns,
            *QualityInspection.__table__.columns,
            *LabAnalysis.__table__.columns,
        ]
        query = query.filter(
            or_(*[cast(column, String).ilike(token) for column in columns])
        )
    rows = query.order_by(CSCase.id.desc()).all()
    quality_ids = [quality.id for _, _, quality, _ in rows if quality]
    memo_rows = (
        db.query(QualityCellMemo)
        .filter(QualityCellMemo.quality_id.in_(quality_ids))
        .all()
        if quality_ids
        else []
    )
    memo_map = {}
    for memo in memo_rows:
        if memo.memo:
            memo_map.setdefault(memo.quality_id, []).append(
                f"{memo.field_name}: {memo.memo}"
            )

    cs_fields = [
        ("date_only", "기준 일자"),
        ("received_at", "접수"),
        ("processed_at", "처리"),
        ("cs_started_at", "상담 시작"),
        ("cs_completed_at", "상담 종료"),
        ("week", "주차"),
        ("transfer_no", "식별코드"),
        ("hospital_name", "병원명"),
        ("rev", "HW Rev"),
        ("seller", "판매처"),
        ("cx_transfer", "CX 이관"),
        ("rnd_transfer", "R&D 이관"),
        ("handover_type", "이관 유형"),
        ("rnd_team", "R&D Team"),
        ("category_main", "유형(대)"),
        ("category_sub", "유형(소)"),
        ("description", "문의사항"),
        ("serial_no", "S/N"),
        ("lot", "LOT"),
        ("size", "Size"),
    ]
    cx_fields = [
        ("cx_yn", "CX 이관"),
        ("rnd_yn", "R&D 이관"),
        ("rnd_team", "R&D Team"),
        ("received_at", "CX 접수"),
        ("processed_at", "CX 처리"),
        ("result_summary", "CX 처리 결과"),
        ("data_loss_type", "데이터 손실"),
        ("data_loss_count", "손실 건수"),
        ("issue1", "불량원인 1"),
        ("depth1_1", "Depth 1-1"),
        ("depth1_2", "Depth 1-2"),
        ("issue2", "불량원인 2"),
        ("depth2_1", "Depth 2-1"),
        ("depth2_2", "Depth 2-2"),
        ("exchange_type", "교품 유형"),
        ("completed_at", "CX 완료일"),
    ]
    exchange_fields = [
        ("exchange_type", "교품 유형"),
        ("return_yn", "회수 여부"),
        ("return_at", "회수일"),
        ("invoice_no", "송장 번호"),
        ("replacement_serial_no", "대치품 S/N"),
    ]
    quality_fields = [
        ("rev", "공정 Rev"),
        ("__hw_rev", "HW Rev"),
        ("return_transfer_date", "회수품 이관일"),
        *QUALITY_COLUMNS,
        ("lab_transfer_date", "품질 분석 완료일"),
    ]
    rnd_fields = [
        (column.name, column.name)
        for column in LabAnalysis.__table__.columns
        if column.name not in {"id", "transfer_no"}
    ]
    headers = (
        ["현재 단계", "진행률", "CX 상태", "교품·회수 상태", "품질 상태", "R&D 상태"]
        + [f"기본/{label}" for _, label in cs_fields]
        + [f"CX/{label}" for _, label in cx_fields]
        + [f"교품·회수/{label}" for _, label in exchange_fields]
        + [f"품질/{label}" for _, label in quality_fields]
        + ["품질/셀 메모"]
        + [f"R&D/{label}" for _, label in rnd_fields]
    )
    wb = Workbook()
    ws = wb.active
    ws.title = "Case Master"
    ws.append(headers)
    for cs, cx, quality, lab in rows:
        state = _case_state(cs, quality, cx, lab)
        values = [
            state["current_stage"],
            state["progress"],
            state["cx_status"],
            state["exchange_status"],
            state["quality_status"],
            state["rnd_status"],
        ]
        values += [
            (
                _display(_hw_rev(cs))
                if field == "rev"
                else _display(getattr(cs, field, None))
            )
            for field, _label in cs_fields
        ]
        values += [
            _display(getattr(cx, field, None)) if cx else "" for field, _ in cx_fields
        ]
        values += [_display(getattr(cs, field, None)) for field, _ in exchange_fields]
        values += [
            (
                _display(_hw_rev(cs))
                if field == "__hw_rev"
                else _display(getattr(quality, field, None)) if quality else ""
            )
            for field, _label in quality_fields
        ]
        values += ["\n".join(memo_map.get(quality.id, [])) if quality else ""]
        values += [
            _display(getattr(lab, field, None)) if lab else ""
            for field, _ in rnd_fields
        ]
        ws.append(values)
    db.close()
    header_fill = PatternFill("solid", fgColor="17324D")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 34
    for index, column in enumerate(ws.columns, 1):
        width = min(
            42,
            max(12, max(len(str(cell.value or "")) for cell in list(column)[:100]) + 2),
        )
        ws.column_dimensions[get_column_letter(index)].width = width
        for cell in column:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": dated_excel_disposition(product, "Master sheet")
        },
    )


@router.get("/{product}/result/excel-legacy")
def download_result_list_excel(product: str):
    db = SessionLocal()
    rows = (
        db.query(CSCase, CXAnalysis, QualityInspection, LabAnalysis)
        .outerjoin(CXAnalysis, CSCase.transfer_no == CXAnalysis.transfer_no)
        .outerjoin(
            QualityInspection, CSCase.transfer_no == QualityInspection.transfer_no
        )
        .outerjoin(LabAnalysis, CSCase.transfer_no == LabAnalysis.transfer_no)
        .filter(
            CSCase.exchange_type.in_(["O", "Cradle"]),
            CSCase.return_at.isnot(None),
            CSCase.product == product,
        )
        .order_by(CSCase.return_at.desc())
        .all()
    )

    wb = Workbook()
    ws = wb.active
    ws.title = "결과 List"
    ws.append(
        [
            "이관번호",
            "접수일",
            "Serial No",
            "CX 원인1",
            "CX Depth1-1",
            "CX Depth1-2",
            "CX 원인2",
            "CX Depth2-1",
            "CX Depth2-2",
            "충전전류(자체)",
            "충전전류(링)",
            "충전전류(크래들)",
            "충전가능",
            "충전로그",
            "배터리(3h)",
            "연결가능",
            "배터리(20분)",
            "신호",
            "BT",
            "외관",
            "방전",
            "방전 결과",
            "침수",
            "캡 높이",
            "불량 내용",
            "품질 판정",
            "품질 분석 완료일",
            "발생 원인",
            "분석 결과",
            "불량 구분",
            "비고",
            "분석 완료일",
        ]
    )

    for cs, cx, quality, lab in rows:
        ws.append(
            [
                cs.transfer_no if cs else "",
                cs.received_at.strftime("%Y-%m-%d") if cs and cs.received_at else "",
                cs.serial_no if cs else "",
                cx.issue1 if cx else "",
                cx.depth1_1 if cx else "",
                cx.depth1_2 if cx else "",
                cx.issue2 if cx else "",
                cx.depth2_1 if cx else "",
                cx.depth2_2 if cx else "",
                quality.current_self_cradle if quality else "",
                quality.current_master_ring if quality else "",
                quality.current_master_cradle if quality else "",
                quality.station_charge if quality else "",
                quality.charge_log if quality else "",
                quality.battery_3h if quality else "",
                quality.disconnect_connectable if quality else "",
                quality.battery_after_disconnect if quality else "",
                quality.signal if quality else "",
                quality.bt if quality else "",
                quality.appearance if quality else "",
                quality.discharge_time if quality else "",
                quality.discharge_result if quality else "",
                quality.water_damage if quality else "",
                quality.cap_height if quality else "",
                quality.defect_detail if quality else "",
                quality.quality_result if quality else "",
                quality.lab_transfer_date if quality else "",
                lab.root_cause if lab else "",
                lab.analysis_result if lab else "",
                lab.defect_type if lab else "",
                lab.remark if lab else "",
                lab.analysis_completed_date if lab else "",
            ]
        )

    db.close()

    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)

    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": dated_excel_disposition(product, "Master sheet")
        },
    )


@router.get("/{product}/api/result/quality/{transfer_no}", response_class=HTMLResponse)
def quality_popup(transfer_no: str, product: str):

    db = SessionLocal()

    quality = (
        db.query(QualityInspection)
        .filter(QualityInspection.transfer_no == transfer_no)
        .first()
    )

    db.close()

    if not quality:
        return "<p>품질 분석 데이터 없음</p>"

    return f"""
    <div class="result-card">

    <table class="result-mini-table">
      <tr><th>충전전류(Cradle)</th><td>{quality.current_self_cradle or '-'}</td></tr>
      <tr><th>충전전류(Master_Ring)</th><td>{quality.current_master_ring or '-'}</td></tr>
      <tr><th>충전전류(Master_Cradle)</th><td>{quality.current_master_cradle or '-'}</td></tr>

      <tr><th>충전 가능 여부</th><td>{quality.station_charge or '-'}</td></tr>
      <tr><th>충전 중 상태(log)</th><td>{quality.charge_log or '-'}</td></tr>
      <tr><th>충전 중 배터리(3h)</th><td>{quality.battery_3h or '-'}</td></tr>

      <tr><th>충전 해제 후 연결</th><td>{quality.disconnect_connectable or '-'}</td></tr>
      <tr><th>충전 해제 후 배터리(20m)</th><td>{quality.battery_after_disconnect or '-'}</td></tr>

      <tr><th>신호</th><td>{quality.signal or '-'}</td></tr>
      <tr><th>BT</th><td>{quality.bt or '-'}</td></tr>
      <tr><th>외관</th><td>{quality.appearance or '-'}</td></tr>

      <tr><th>방전</th><td>{quality.discharge_time or '-'}</td></tr>
      <tr><th>방전 후 상태</th><td>{quality.discharge_result or '-'}</td></tr>

      <tr><th>침수</th><td>{quality.water_damage or '-'}</td></tr>
      <tr><th>캡높이</th><td>{quality.cap_height or '-'}</td></tr>
      <tr><th>불량내용</th><td>{quality.defect_detail or '-'}</td><tr>
      <tr><th>품질 판정</th><td><b>{quality.quality_result or '-'}</b></td></tr>

      <tr><th>품질 분석 완료일</th><td>{quality.lab_transfer_date or '-'}</td></tr>
    </table>

    </div>
    """


@router.get("/{product}/api/result/rnd/{transfer_no}", response_class=HTMLResponse)
def rnd_popup(transfer_no: str, product: str):

    db = SessionLocal()

    cs, lab = (
        db.query(CSCase, LabAnalysis)
        .outerjoin(LabAnalysis, CSCase.transfer_no == LabAnalysis.transfer_no)
        .filter(CSCase.transfer_no == transfer_no, CSCase.product == product)
        .first()
    )

    db.close()

    if not lab:
        return "<p>RND 분석 데이터 없음</p>"

    return f"""
    <div class="result-card">

    <table class="result-mini-table">

      <tr>
        <th>Serial No</th>
        <td>{lab.serial_no or '-'}</td>
      </tr>

      <tr>
        <th>분석 결과</th>
        <td>{lab.analysis_result or '-'}</td>
      </tr>

      <tr>
        <th>불량 구분</th>
        <td>{lab.defect_type or '-'}</td>
      </tr>

      <tr>
        <th>비고</th>
        <td>{lab.remark or '-'}</td>
      </tr>

      <tr>
        <th>분석 완료일</th>
        <td>{lab.analysis_completed_date or '-'}</td>
      </tr>

    </table>
    """